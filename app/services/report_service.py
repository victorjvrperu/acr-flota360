from sqlalchemy.orm import Session
from sqlalchemy import func
from datetime import datetime, date
from typing import Optional
import urllib.parse
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.heavytrack import (
    Proyecto, Personal, MaquinariaEquipo, TurnoOperacion,
    ControlCombustible, OrdenMantenimiento, InventarioRepuesto
)
from app.services.heavytrack_service import HeavyTrackService


class ReportService:

    @staticmethod
    def generar_mensaje_whatsapp_gerencia(db: Session, telefono_destino: Optional[str] = None) -> dict:
        """
        Construye el resumen ejecutivo formateado para envío por WhatsApp a Gerencia
        con enlaces al reporte detallado y KPIs clave.
        """
        from app.models.heavytrack import AjusteSistema
        if not telefono_destino or not telefono_destino.strip():
            aj = db.query(AjusteSistema).filter(AjusteSistema.clave == "whatsapp_gerencia_telefono").first()
            telefono_destino = aj.valor if aj and aj.valor else "+51998887766"

        resumen = HeavyTrackService.obtener_resumen_dashboard(db)
        hoy_str = datetime.utcnow().strftime("%d/%m/%Y %H:%M")

        # Equipos con alertas de semáforo
        equipos = db.query(MaquinariaEquipo).all()
        alertas_txt = []
        for eq in equipos:
            color, rest = HeavyTrackService.calcular_semaforo_equipo(eq)
            tipo_c = (eq.tipo_control or "HOROMETRO").upper()
            unit = "km" if tipo_c == "KILOMETRAJE" else "h"
            nom_eq = eq.nombre or f"{eq.marca or ''} {eq.modelo or ''}".strip() or "Equipo"
            if color == "rojo":
                alertas_txt.append(f"  🔴 *{eq.codigo_interno}* ({nom_eq}): Mantenimiento *VENCIDO* hace {abs(rest):.1f} {unit}.")
            elif color == "amarillo":
                alertas_txt.append(f"  🟡 *{eq.codigo_interno}* ({nom_eq}): Próximo a servicio. Faltan *{rest:.1f} {unit}*.")

        alertas_str = "\n".join(alertas_txt) if alertas_txt else "  ✅ Ninguna alerta crítica. Flota 100% al día."

        # Proyectos
        proyectos_txt = []
        for p in resumen["horas_por_proyecto"]:
            proyectos_txt.append(f"  • *{p['codigo']}* ({p['nombre']}): {p['horas']} horas.")
        proyectos_str = "\n".join(proyectos_txt) if proyectos_txt else "  • En Base / Taller Central"

        mensaje = (
            f"🚜 *ACR FLOTA 360° — INFORME EJECUTIVO*\n"
            f"📅 *Corte:* {hoy_str}\n\n"
            f"📊 *KPIs GLOBALES DE FLOTA:*\n"
            f"• *Horas Totales Trabajadas:* {resumen['horas_totales_trabajadas']} h\n"
            f"• *Equipos Operativos:* {resumen['semaforo_flota']['verde']} de {resumen['total_equipos']}\n"
            f"• *Equipos en Mantenimiento:* {resumen['equipos_en_taller']} unidades\n"
            f"• *Combustible Mes:* {resumen['galones_combustible_mes']} Gal (S/ {resumen['costo_combustible_mes']:,.2f})\n\n"
            f"🚦 *ESTADO DE SEMÁFOROS Y MANTENIMIENTO:*\n"
            f"{alertas_str}\n\n"
            f"🏗️ *ACTIVIDAD POR FRENTE DE OBRA:*\n"
            f"{proyectos_str}\n\n"
            f"📄 *Reporte Imprimible / PDF:* http://100.84.247.4:8000/api/v1/heavy/reportes/pdf\n"
            f"🔒 *Plataforma:* ACR Flota 360° Cloud"
        )

        # Enlace directo de WhatsApp Web / App
        tel_clean = telefono_destino.replace('+', '').replace(' ', '').replace('-', '')
        encoded_msg = urllib.parse.quote(mensaje)
        whatsapp_url = f"https://api.whatsapp.com/send?phone={tel_clean}&text={encoded_msg}"

        return {
            "mensaje": mensaje,
            "mensaje_texto": mensaje,
            "telefono_destino": telefono_destino,
            "telefono_clean": tel_clean,
            "whatsapp_link": whatsapp_url,
            "timestamp": hoy_str
        }

    @staticmethod
    def generar_html_reporte_pdf(db: Session) -> str:
        """
        Genera el documento imprimible/PDF para entrega a Encargados de Área y Gerencia.
        """
        resumen = HeavyTrackService.obtener_resumen_dashboard(db)
        equipos = db.query(MaquinariaEquipo).all()
        proyectos = db.query(Proyecto).all()
        repuestos_bajos = db.query(InventarioRepuesto).filter(
            InventarioRepuesto.stock_actual <= InventarioRepuesto.stock_minimo
        ).all()

        fecha_str = datetime.utcnow().strftime("%d/%m/%Y")

        filas_equipos = ""
        for eq in equipos:
            color, horas_rest = HeavyTrackService.calcular_semaforo_equipo(eq)
            color_badge = "#10B981" if color == "verde" else ("#F59E0B" if color == "amarillo" else "#EF4444")
            estado_txt = "OPERATIVO" if color == "verde" else ("<25h PRÓX" if color == "amarillo" else "VENCIDO")

            filas_equipos += f"""
            <tr>
              <td style="font-weight:bold; color:#FFB800;">{eq.codigo_interno}</td>
              <td>{eq.marca} {eq.modelo}</td>
              <td>{eq.tipo.upper()}</td>
              <td style="text-align:right;">{eq.horometro_actual:.1f} h</td>
              <td style="text-align:right;">{eq.horometro_proximo_servicio:.1f} h</td>
              <td style="text-align:right; font-weight:bold; color:{color_badge};">{horas_rest:.1f} h</td>
              <td style="text-align:center;"><span style="background:{color_badge}; color:#000; padding:2px 8px; border-radius:4px; font-weight:bold; font-size:10px;">{estado_txt}</span></td>
            </tr>
            """

        filas_repuestos = ""
        for r in repuestos_bajos:
            filas_repuestos += f"""
            <tr>
              <td><strong>{r.codigo_parte}</strong></td>
              <td>{r.nombre}</td>
              <td style="text-align:center; color:#EF4444; font-weight:bold;">{r.stock_actual} (Mín: {r.stock_minimo})</td>
              <td style="text-align:right;">S/ {r.precio_unitario:.2f}</td>
              <td>{r.ubicacion_almacen or 'Almacén Central'}</td>
            </tr>
            """

        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
          <meta charset="utf-8">
          <title>Reporte Ejecutivo HeavyTrack - {fecha_str}</title>
          <style>
            body {{ font-family: Arial, sans-serif; background: #fff; color: #111; margin: 30px; font-size: 12px; }}
            .header {{ border-bottom: 2px solid #FFB800; padding-bottom: 15px; margin-bottom: 20px; display: flex; justify-content: space-between; align-items: center; }}
            .title {{ font-size: 20px; font-weight: bold; color: #121417; }}
            .badge-brand {{ background: #FFB800; color: #121417; padding: 4px 10px; font-weight: bold; border-radius: 4px; display: inline-block; }}
            .kpi-box {{ display: inline-block; width: 23%; background: #f4f4f5; padding: 12px; border-radius: 6px; margin-right: 1.5%; box-sizing: border-box; }}
            .kpi-val {{ font-size: 18px; font-weight: bold; color: #121417; margin-top: 4px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; }}
            th {{ background: #1E2228; color: #fff; text-align: left; padding: 8px; font-size: 11px; text-transform: uppercase; }}
            td {{ padding: 8px; border-bottom: 1px solid #e5e7eb; }}
            .section-title {{ font-size: 14px; font-weight: bold; margin-top: 25px; color: #1E2228; border-left: 4px solid #FFB800; padding-left: 8px; }}
            .footer {{ margin-top: 40px; font-size: 10px; color: #6b7280; text-align: center; border-top: 1px solid #e5e7eb; padding-top: 10px; }}
            @media print {{ body {{ margin: 0; }} }}
          </style>
        </head>
        <body>
          <div class="header">
            <div>
              <div class="badge-brand">HEAVYTRACK ENTERPRISE</div>
              <div class="title" style="margin-top:6px;">REPORTE INTEGRAL DE FLOTA Y OPERACIONES</div>
              <div style="color: #6b7280; font-size: 11px;">Documento Oficial para Gerencia y Jefaturas de Obra</div>
            </div>
            <div style="text-align: right; font-size: 11px;">
              <strong>Fecha de Emisión:</strong> {fecha_str}<br>
              <strong>Sistema:</strong> HeavyTrack On-Premise v2.0
            </div>
          </div>

          <!-- KPIs -->
          <div style="margin-bottom: 20px;">
            <div class="kpi-box">
              <span style="color:#6b7280; font-size:10px; font-weight:bold;">HORAS TRABAJADAS</span>
              <div class="kpi-val">{resumen['horas_totales_trabajadas']:.1f} h</div>
            </div>
            <div class="kpi-box">
              <span style="color:#6b7280; font-size:10px; font-weight:bold;">EQUIPOS OPERATIVOS</span>
              <div class="kpi-val" style="color:#10B981;">{resumen['semaforo_flota']['verde']} / {resumen['total_equipos']}</div>
            </div>
            <div class="kpi-box">
              <span style="color:#6b7280; font-size:10px; font-weight:bold;">ALERTAS &lt; 25h</span>
              <div class="kpi-val" style="color:#EF4444;">{resumen['alertas_mantenimiento_25h']} Críticas</div>
            </div>
            <div class="kpi-box" style="margin-right:0;">
              <span style="color:#6b7280; font-size:10px; font-weight:bold;">COMBUSTIBLE MES</span>
              <div class="kpi-val">{resumen['galones_combustible_mes']:.1f} Gal</div>
            </div>
          </div>

          <!-- Tabla Maquinaria -->
          <div class="section-title">ESTADO DE MAQUINARIA PESADA Y HORÓMETROS (INTERVALO 250H)</div>
          <table>
            <thead>
              <tr>
                <th>Código</th>
                <th>Equipo / Modelo</th>
                <th>Tipo</th>
                <th style="text-align:right;">Horómetro</th>
                <th style="text-align:right;">Próx. Servicio</th>
                <th style="text-align:right;">Horas Restantes</th>
                <th style="text-align:center;">Semáforo</th>
              </tr>
            </thead>
            <tbody>
              {filas_equipos}
            </tbody>
          </table>

          <!-- Repuestos Críticos -->
          <div class="section-title">REPUESTOS CRÍTICOS POR REABASTECER (STOCK &le; MÍNIMO)</div>
          <table>
            <thead>
              <tr>
                <th>Código Parte</th>
                <th>Descripción</th>
                <th style="text-align:center;">Stock Actual / Mínimo</th>
                <th style="text-align:right;">Precio Unit.</th>
                <th>Ubicación</th>
              </tr>
            </thead>
            <tbody>
              {filas_repuestos if filas_repuestos else '<tr><td colspan="5" style="text-align:center; color:#10B981;">Stock en niveles óptimos.</td></tr>'}
            </tbody>
          </table>

          <div class="footer">
            HeavyTrackEnterprise © 2026 — Generado automáticamente con FastAPI y SQLAlchemy. Firma y sello de supervisión de área.
          </div>
        </body>
        </html>
        """
        return html

    @staticmethod
    def generar_excel_control_horometros(
        db: Session,
        proyecto_id: Optional[int] = None,
        equipo_id: Optional[int] = None,
        operador_id: Optional[int] = None,
        fecha_desde: Optional[date] = None,
        fecha_hasta: Optional[date] = None
    ) -> io.BytesIO:
        """
        Genera un archivo Excel profesional (.xlsx) con el Parte Diario
        de Control de Horómetros, Kilometraje, Combustible y Horas Efectivas.
        """
        query = db.query(TurnoOperacion)
        if proyecto_id:
            query = query.filter(TurnoOperacion.proyecto_id == proyecto_id)
        if equipo_id:
            query = query.filter(TurnoOperacion.equipo_id == equipo_id)
        if operador_id:
            query = query.filter(TurnoOperacion.operador_id == operador_id)
        if fecha_desde:
            query = query.filter(TurnoOperacion.fecha >= fecha_desde)
        if fecha_hasta:
            query = query.filter(TurnoOperacion.fecha <= fecha_hasta)

        turnos = query.order_by(TurnoOperacion.fecha.desc(), TurnoOperacion.id.desc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Parte Diario Horometros"
        ws.views.sheetView[0].showGridLines = True

        # Paleta Industrial ACR FLOTA 360
        font_titulo = Font(name="Calibri", size=13, bold=True, color="01A95B")
        font_subtitulo = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_meta = Font(name="Calibri", size=9, italic=True, color="4B5563")
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10, color="1F2937")
        font_data_bold = Font(name="Calibri", size=10, bold=True, color="01A95B")
        font_total = Font(name="Calibri", size=10, bold=True, color="111827")

        fill_brand = PatternFill(start_color="121417", end_color="121417", fill_type="solid")
        fill_sub = PatternFill(start_color="1E2228", end_color="1E2228", fill_type="solid")
        fill_header = PatternFill(start_color="282D35", end_color="282D35", fill_type="solid")
        fill_zebra_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        fill_zebra_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_totales = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

        thin_side = Side(border_style="thin", color="D1D5DB")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        border_top_thick = Border(
            left=thin_side, right=thin_side,
            top=Side(border_style="medium", color="01A95B"),
            bottom=Side(border_style="double", color="111827")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Fila 1: Título Marca
        ws.merge_cells("A1:N1")
        cell_t1 = ws["A1"]
        cell_t1.value = "ACR FLOTA 360° — SISTEMA DE GESTIÓN Y CONTROL DE MAQUINARIA PESADA"
        cell_t1.font = font_titulo
        cell_t1.fill = fill_brand
        cell_t1.alignment = align_center
        ws.row_dimensions[1].height = 26

        # Fila 2: Subtítulo de Reporte
        ws.merge_cells("A2:N2")
        cell_t2 = ws["A2"]
        cell_t2.value = "PARTE DIARIO DE CONTROL DE HORÓMETROS, KILOMETRAJE Y COMBUSTIBLE"
        cell_t2.font = font_subtitulo
        cell_t2.fill = fill_sub
        cell_t2.alignment = align_center
        ws.row_dimensions[2].height = 20

        # Fila 3: Metadatos y fecha de corte
        hoy_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        filtro_txt = []
        if fecha_desde: filtro_txt.append(f"Desde: {fecha_desde.strftime('%d/%m/%Y')}")
        if fecha_hasta: filtro_txt.append(f"Hasta: {fecha_hasta.strftime('%d/%m/%Y')}")
        filtro_info = " | ".join(filtro_txt) if filtro_txt else "Rango: Todos los registros"

        ws.merge_cells("A3:N3")
        cell_t3 = ws["A3"]
        cell_t3.value = f"Fecha de Exportación: {hoy_str}  |  {filtro_info}  |  Total Registros: {len(turnos)}"
        cell_t3.font = font_meta
        cell_t3.alignment = align_left
        ws.row_dimensions[3].height = 18

        # Fila 4: Espacio
        ws.row_dimensions[4].height = 6

        # Fila 5: Encabezados de Columna
        headers = [
            "N°", "FECHA", "FRENTE / OBRA", "CÓD. EQUIPO", "MAQUINARIA (MODELO)",
            "OPERADOR A CARGO", "DNI", "HORÓMETRO INICIO", "HORÓMETRO FIN",
            "Δ HORAS TRAB.", "COMBUSTIBLE (GAL)", "RATIO (GAL/H)", "INCIDENCIA / NOTA", "ESTADO"
        ]
        ws.row_dimensions[5].height = 26

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=header)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_header
            cell.border = border_cell

        # Filas 6..N: Datos
        row_idx = 6
        total_horas = 0.0
        total_galones = 0.0

        for idx, t in enumerate(turnos, start=1):
            ws.row_dimensions[row_idx].height = 20
            row_fill = fill_zebra_even if idx % 2 != 0 else fill_zebra_odd

            frente_nom = t.proyecto.nombre if t.proyecto else "Sin Asignar"
            cod_eq = t.equipo.codigo_interno if t.equipo else "N/A"
            eq_nom = f"{t.equipo.marca} {t.equipo.modelo}" if t.equipo else "N/A"
            op_nom = f"{t.operador.nombres} {t.operador.apellidos}" if t.operador else "N/A"
            op_dni = t.operador.dni if t.operador else "N/A"

            h_ini = t.horometro_inicio if t.horometro_inicio is not None else 0.0
            h_fin = t.horometro_fin if t.horometro_fin is not None else 0.0
            h_trab = t.horas_trabajadas if t.horas_trabajadas is not None else 0.0
            gal = t.combustible_galones if t.combustible_galones is not None else 0.0

            ratio = round(gal / h_trab, 2) if (gal > 0 and h_trab > 0) else 0.0
            incidencia = t.incidencia_descripcion if t.hubo_incidencia else "Conforme"

            total_horas += h_trab
            total_galones += gal

            valores_fila = [
                (idx, align_center, False, "@"),
                (t.fecha.strftime("%d/%m/%Y"), align_center, False, "@"),
                (frente_nom, align_left, False, "@"),
                (cod_eq, align_center, True, "@"),
                (eq_nom, align_left, False, "@"),
                (op_nom, align_left, False, "@"),
                (op_dni, align_center, False, "@"),
                (h_ini, align_right, False, '0.0 "h"'),
                (h_fin if h_fin > 0 else "-", align_right, False, '0.0 "h"' if h_fin > 0 else "@"),
                (h_trab, align_right, True, '0.0 "h"'),
                (gal if gal > 0 else 0.0, align_right, False, '0.0 "Gal"'),
                (ratio if ratio > 0 else "-", align_right, False, '0.00 "Gal/h"' if ratio > 0 else "@"),
                (incidencia, align_left, False, "@"),
                (t.estado_sync or "SINCRONIZADO", align_center, False, "@")
            ]

            for col_idx, (val, alignment, is_bold, num_format) in enumerate(valores_fila, start=1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.fill = row_fill
                c.border = border_cell
                c.alignment = alignment
                c.font = font_data_bold if is_bold else font_data
                c.number_format = num_format

            row_idx += 1

        # Fila de Totales
        ws.row_dimensions[row_idx].height = 24
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=9)
        lbl_tot = ws.cell(row=row_idx, column=1, value="TOTALES Y PROMEDIOS ACUMULADOS:")
        lbl_tot.font = font_total
        lbl_tot.fill = fill_totales
        lbl_tot.alignment = align_right

        for c_idx in range(1, 10):
            ws.cell(row=row_idx, column=c_idx).border = border_top_thick
            ws.cell(row=row_idx, column=c_idx).fill = fill_totales

        # Horas totales
        c_tot_h = ws.cell(row=row_idx, column=10, value=round(total_horas, 1))
        c_tot_h.font = font_total
        c_tot_h.fill = fill_totales
        c_tot_h.border = border_top_thick
        c_tot_h.alignment = align_right
        c_tot_h.number_format = '0.0 "h"'

        # Galones totales
        c_tot_gal = ws.cell(row=row_idx, column=11, value=round(total_galones, 1))
        c_tot_gal.font = font_total
        c_tot_gal.fill = fill_totales
        c_tot_gal.border = border_top_thick
        c_tot_gal.alignment = align_right
        c_tot_gal.number_format = '0.0 "Gal"'

        # Ratio promedio
        ratio_prom = round(total_galones / total_horas, 2) if total_horas > 0 else 0.0
        c_ratio = ws.cell(row=row_idx, column=12, value=f"{ratio_prom:.2f} Gal/h" if ratio_prom > 0 else "-")
        c_ratio.font = font_total
        c_ratio.fill = fill_totales
        c_ratio.border = border_top_thick
        c_ratio.alignment = align_right

        for c_idx in [13, 14]:
            c_empty = ws.cell(row=row_idx, column=c_idx, value="")
            c_empty.fill = fill_totales
            c_empty.border = border_top_thick

        # Auto-ajuste de anchos de columna con márgenes óptimos
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row in [1, 2, 3] or (cell.row == row_idx and col[0].column < 10):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
